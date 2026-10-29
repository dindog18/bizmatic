import asyncio
from pathlib import Path

import pandas as pd
import pymupdf
from docx import Document

from ..schemas.files import FileType, StoredFileMetadata
from ..utils.file_types import identify_file_type_metadata
from ..utils.markdown import markdown_table


async def document_file_text_extraction(stored_file: StoredFileMetadata) -> str:
    """Extract a supported document and return its contents as Markdown.

    Images are intentionally handled by the separate image OCR pipeline.
    """
    file_type: FileType = identify_file_type_metadata(stored_file)
    # Images should go through image_file_ocr_text_extraction instead.
    if file_type == "image":
        raise ValueError("Expected a document file, not an image file.")
    if file_type == "unknown":
        raise ValueError(f"Unsupported or unknown file type: {stored_file.original_file_name}")

    file_path = Path(stored_file.storage_path)
    if not file_path.exists():
        raise FileNotFoundError(f"Stored file does not exist: {file_path}")

    # Keep each format implementation separate while sharing dispatch here.
    extractors = {
        "pdf": extract_pdf_to_markdown,
        "word": extract_word_to_markdown,
        "excel": extract_excel_to_markdown,
        "text": extract_text_to_markdown,
    }
    extractor = extractors.get(file_type)
    if extractor is None:
        raise ValueError(f"Unsupported document file type: {file_type}")
    return await asyncio.to_thread(extractor, file_path)


def extract_pdf_to_markdown(file_path: Path) -> str:
    """Preserve PDF tables and reading order instead of flattening every cell."""
    parts = []
    with pymupdf.open(file_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            tables = page.find_tables().tables
            regions = []
            elements = []
            for table in tables:
                # Merged cells can cause unrelated sections to be detected as
                # one table. Split where the actual column layout changes.
                group, group_cells, previous_columns, top = [], [], None, table.bbox[1]
                def emit_group():
                    # Short metadata/merged-label sections read better as
                    # positioned text; reserve tables for actual row grids.
                    if len(group) >= 2 and len(group[0]) >= 3 and any(any(cell for cell in row) for row in group):
                        elements.append((top, table.bbox[0], _pdf_table(group)))
                        regions.extend(group_cells)
                for row, geometry in zip(table.extract(), table.rows):
                    columns = tuple(i for i, cell in enumerate(row) if cell is not None)
                    if not columns:
                        continue
                    if group and columns != previous_columns:
                        emit_group()
                        group, group_cells = [], []
                    if not group:
                        top = geometry.bbox[1]
                    group.append([str(row[i] or '').strip() for i in columns])
                    group_cells.extend(pymupdf.Rect(geometry.cells[i]) for i in columns if geometry.cells[i] is not None)
                    previous_columns = columns
                if group:
                    emit_group()
            # Keep all text outside tables, including titles and footnotes.
            # Group words by their visual baseline to retain label/value order.
            lines = []
            for word in sorted(page.get_text('words'), key=lambda w: (w[1], w[0])):
                center = pymupdf.Point((word[0]+word[2])/2, (word[1]+word[3])/2)
                if any(center in region for region in regions):
                    continue
                if not lines or abs(word[1]-lines[-1][0]) > 3:
                    lines.append((word[1], []))
                lines[-1][1].append(word)
            for y, words in lines:
                words.sort(key=lambda w: w[0])
                elements.append((y, words[0][0], ' '.join(w[4] for w in words)))
            content = '\n\n'.join(value for _, _, value in sorted(elements) if value.strip())
            if not content.strip():
                # A genuinely blank rendered page can be skipped; a page with
                # visible image or drawing content still needs OCR.
                scale = min(1.0, 1024 / max(page.rect.width, page.rect.height, 1))
                pixels = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).samples
                visible_drawing = any(
                    drawing.get(name) is not None and any(channel < 1 for channel in drawing[name])
                    for drawing in page.get_drawings() for name in ('color', 'fill')
                )
                if not page.get_image_info() and not visible_drawing and pixels.count(255) == len(pixels):
                    continue
                raise ValueError(f'Page {page_number} has no extractable text. Use OCR for scanned PDFs.')
            parts.append(f'## Page {page_number}\n\n{content}')
    if not parts:
        raise ValueError('This PDF has no extractable text. Use OCR for scanned PDFs.')
    return '\n\n'.join(parts).strip()


def _pdf_table(rows):
    # Keep one-row metadata sections as data rather than column headings.
    # Multi-row grids retain their first row as the document's table header.
    if len(rows) == 1:
        rows = [[''] * len(rows[0]), *rows]
    return markdown_table(rows)


def extract_word_to_markdown(file_path: Path) -> str:
    """Convert a .docx file to basic Markdown, including headings and tables."""
    if file_path.suffix.lower() == ".doc":
        raise ValueError("Legacy .doc files are not directly supported. Convert the file to .docx first.")
    document = Document(str(file_path))
    parts = []
    # Preserve heading hierarchy and regular paragraph text.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style
        style_name = (style.name or "").lower() if style is not None else ""
        if style_name.startswith("heading"):
            try:
                level = int(style_name.split()[-1])
            except (ValueError, IndexError):
                level = 2
            parts.append(f"{'#' * min(max(level, 1), 6)} {text}")
        else:
            parts.append(text)
    # Append document tables after the paragraphs.
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            parts.append(markdown_table(rows))
    return "\n\n".join(parts).strip()


def extract_excel_to_markdown(file_path: Path) -> str:
    """Convert each worksheet in an Excel workbook into a Markdown table."""
    workbook = pd.ExcelFile(file_path)
    parts = []
    for sheet_name in workbook.sheet_names:
        # Read values as strings and render missing cells as blanks.
        dataframe = pd.read_excel(workbook, sheet_name=sheet_name, dtype=str).fillna("")
        parts.append(f"## Sheet: {sheet_name}")
        parts.append("_Empty worksheet_" if dataframe.empty else dataframe.to_markdown(index=False))
    return "\n\n".join(parts).strip()


def extract_text_to_markdown(file_path: Path) -> str:
    """Convert TXT and CSV files to Markdown (CSV becomes a table)."""
    if file_path.suffix.lower() == ".csv":
        return extract_csv_to_markdown(file_path)
    return file_path.read_text(encoding="utf-8", errors="replace").strip()


def extract_csv_to_markdown(file_path: Path) -> str:
    """Convert a CSV file into a Markdown table."""
    # Read values as strings and render missing cells as blanks.
    # PyCharm cannot resolve pandas' generic Unpack overloads for this valid call.
    # noinspection PyArgumentList
    dataframe = pd.read_csv(str(file_path), dtype=str).fillna("")
    return "_Empty CSV file_" if dataframe.empty else dataframe.to_markdown(index=False)
