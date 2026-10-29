"""Synthetic hostile documents must be rejected before expensive format parsing."""
import asyncio
import struct
import zlib
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import pymupdf
from docx import Document
from openpyxl import Workbook
from django.test import SimpleTestCase

from document_extraction.schemas.files import StoredFileMetadata
from document_extraction.services import documents, parser_limits as limits


class ParserLimitTests(SimpleTestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)

    def archive(self, members, name='document.docx', compression=ZIP_DEFLATED):
        path = self.folder / name
        with ZipFile(path, 'w', compression=compression) as archive:
            for key, value in members:
                archive.writestr(key, value)
        return path

    def test_normal_docx_and_xlsx_still_extract(self):
        word = self.folder / 'normal.docx'
        doc = Document()
        doc.add_paragraph('Invoice 100')
        doc.save(word)
        self.assertIn('Invoice 100', documents.extract_word_to_markdown(word))
        excel = self.folder / 'renamed.xls'
        book = Workbook()
        book.active.append(['Order', 'Amount'])
        book.active.append(['PO-001', 100])
        book.save(excel)
        self.assertIn('PO-001', documents.extract_excel_to_markdown(excel))

    def test_member_total_and_count_limits_precede_office_parsers(self):
        cases = [({'MAX_MEMBER_BYTES': 100}, [('media', b'A' * 101)]),
                 ({'MAX_EXPANDED_BYTES': 100}, [('one', b'A' * 60), ('two', b'B' * 60)]),
                 ({'MAX_ARCHIVE_MEMBERS': 1}, [('one', b'A'), ('two', b'B')])]
        for overrides, members in cases:
            for name, extractor, target in [('bomb.docx', documents.extract_word_to_markdown, 'Document'),
                                             ('bomb.xls', documents.extract_excel_to_markdown, 'pd.ExcelFile')]:
                with self.subTest(overrides=overrides, name=name):
                    path = self.archive(members, name)
                    with patch.multiple(limits, **overrides), patch.object(documents, 'Document') as word, patch('document_extraction.services.documents.pd.ExcelFile') as excel:
                        with self.assertRaises(ValueError):
                            extractor(path)
                        word.assert_not_called()
                        excel.assert_not_called()

    def test_forged_small_member_size_cannot_hide_large_deflate_output(self):
        payload = b'A' * (1024 * 1024)
        path = self.archive([('word/media/image.bin', payload)])
        raw = bytearray(path.read_bytes())
        central = raw.index(b'PK\x01\x02')
        # Craft internally consistent metadata and CRC for a truncated prefix.
        # Unbounded ZipFile.read would inflate the full MiB, then return 100 bytes.
        crc = zlib.crc32(payload[:100])
        struct.pack_into('<I', raw, 14, crc)
        struct.pack_into('<I', raw, 22, 100)
        struct.pack_into('<I', raw, central + 16, crc)
        struct.pack_into('<I', raw, central + 24, 100)
        path.write_bytes(raw)
        original_decoder = zlib.decompressobj
        output_sizes = []

        class ObservedDecoder:
            def __init__(self, *args):
                self.decoder = original_decoder(*args)

            def __getattr__(self, name):
                return getattr(self.decoder, name)

            def decompress(self, data, max_length):
                self_output = self.decoder.decompress(data, max_length)
                output_sizes.append(len(self_output))
                return self_output

        with patch.object(documents, 'Document') as parser, patch.object(limits.zlib, 'decompressobj', ObservedDecoder):
            with self.assertRaisesMessage(ValueError, 'size does not match'):
                documents.extract_word_to_markdown(path)
            parser.assert_not_called()
        self.assertLessEqual(max(output_sizes), 101)

    def test_valid_multi_chunk_stored_and_deflated_parts_pass(self):
        for compression in (ZIP_STORED, ZIP_DEFLATED):
            with self.subTest(compression=compression):
                path = self.archive([('large-part', b'A' * (3 * limits.CHUNK_BYTES + 19))], compression=compression)
                limits.check_office_archive(path)

    def test_directory_count_is_checked_before_indexing_even_if_end_record_lies(self):
        path = self.archive([(str(index), b'') for index in range(6)])
        raw = bytearray(path.read_bytes())
        end = raw.rindex(b'PK\x05\x06')
        struct.pack_into('<HH', raw, end + 8, 1, 1)
        path.write_bytes(raw)
        with patch.object(limits, 'MAX_ARCHIVE_MEMBERS', 4), patch.object(limits, 'ZipFile') as parser:
            with self.assertRaisesMessage(ValueError, 'too many entries'):
                limits.check_office_archive(path)
            parser.assert_not_called()

    def test_small_zip64_package_passes(self):
        with patch.object(zipfile, 'ZIP64_LIMIT', 1):
            path = self.archive([('part', b'ordinary content')])
        limits.check_office_archive(path)

    def test_corrupt_stream_checksum_encryption_and_unsupported_compression_rejected(self):
        for mutation in ('crc', 'deflate', 'encrypted', 'compression'):
            with self.subTest(mutation=mutation):
                path = self.archive([('part', b'A' * 1000)])
                raw = bytearray(path.read_bytes())
                central = raw.index(b'PK\x01\x02')
                if mutation == 'crc':
                    struct.pack_into('<I', raw, 14, 0)
                    struct.pack_into('<I', raw, central + 16, 0)
                elif mutation == 'deflate':
                    raw[30 + len('part')] = 0xff
                elif mutation == 'encrypted':
                    struct.pack_into('<H', raw, 6, 1)
                    struct.pack_into('<H', raw, central + 8, 1)
                else:
                    struct.pack_into('<H', raw, 8, 99)
                    struct.pack_into('<H', raw, central + 10, 99)
                path.write_bytes(raw)
                with self.assertRaises(ValueError):
                    limits.check_office_archive(path)

    def test_size_limit_uses_actual_file_for_direct_and_async_calls(self):
        path = self.archive([('part', b'hello')])
        metadata = StoredFileMetadata(original_file_name=path.name, stored_file_name=path.name,
            storage_path=str(path), mime_type=None, file_size=1)
        with patch.object(limits, 'MAX_FILE_BYTES', 10), patch.object(documents, 'Document') as parser:
            with self.assertRaisesMessage(ValueError, 'extraction limit'):
                asyncio.run(documents.document_file_text_extraction(metadata))
            parser.assert_not_called()

    def grid(self, name='grid.pdf', *, lines=160, rectangles=False, rotation=0):
        path = self.folder / name
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            shape = page.new_shape()
            for index in range(lines):
                point = 40 + index * 2
                if rectangles:
                    shape.draw_rect(pymupdf.Rect(point, 40, point + 2, 400))
                else:
                    shape.draw_line((point, 40), (point, 400))
                    shape.draw_line((40, point), (400, point))
            shape.finish(color=(0, 0, 0))
            shape.commit()
            page.insert_text((42, 52), 'Invoice 100')
            page.set_rotation(rotation)
            pdf.save(path, deflate=True)
        return path

    def test_dense_single_path_and_rectangles_rejected_before_table_detection(self):
        for rectangles in (False, True):
            path = self.grid(rectangles=rectangles)
            with patch.object(pymupdf.Page, 'find_tables') as parser:
                with self.assertRaisesMessage(ValueError, 'table-complexity'):
                    documents.extract_pdf_to_markdown(path)
                parser.assert_not_called()

    def test_page_and_cumulative_limits(self):
        path = self.folder / 'pages.pdf'
        with pymupdf.open() as pdf:
            for _ in range(3):
                pdf.new_page().insert_text((72, 72), 'Invoice 100')
            pdf.save(path)
        with patch.object(limits, 'MAX_PDF_PAGES', 2), patch.object(pymupdf.Page, 'find_tables') as parser:
            with self.assertRaisesMessage(ValueError, 'page extraction limit'):
                documents.extract_pdf_to_markdown(path)
            parser.assert_not_called()
        path = self.grid(lines=3)
        with patch.object(limits, 'MAX_DOCUMENT_EDGES', 1), patch.object(pymupdf.Page, 'find_tables') as parser:
            with self.assertRaisesMessage(ValueError, 'document table-complexity'):
                documents.extract_pdf_to_markdown(path)
            parser.assert_not_called()

    def test_ordinary_table_and_rotated_table_preserve_cells(self):
        for rotation in (0, 90):
            path = self.folder / 'table.pdf'
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                shape = page.new_shape()
                for x in (40, 140, 240, 340):
                    shape.draw_line((x, 40), (x, 160))
                for y in (40, 80, 120, 160):
                    shape.draw_line((40, y), (340, y))
                shape.finish(color=(0, 0, 0)); shape.commit()
                for x, value in ((45, 'Order'), (145, 'Item'), (245, 'Amount')):
                    page.insert_text((x, 60), value)
                for x, value in ((45, 'PO-001'), (145, 'Widget'), (245, '100')):
                    page.insert_text((x, 100), value)
                for x, value in ((45, 'PO-002'), (145, 'Gadget'), (245, '200')):
                    page.insert_text((x, 140), value)
                page.set_rotation(rotation)
                pdf.save(path)
            result = documents.extract_pdf_to_markdown(path)
            self.assertIn('PO-001', result)
            self.assertIn('Widget', result)
            self.assertIn('100', result)
            self.assertIn('|', result)
